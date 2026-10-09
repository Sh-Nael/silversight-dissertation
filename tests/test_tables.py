"""The table pack: formatting only, and faithful to the report files."""

import numpy as np
import pandas as pd
import pytest

from xaglab.eval import tables as T
from xaglab.paths import REPORTS_DIR


def test_cell_formats():
    assert T.pval(0.0004) == "<0.001*"
    assert T.pval(0.049) == "0.049*"
    assert T.pval(0.05) == "0.050"
    assert T.pval(np.nan) == "–"
    assert T.pct(0.5382) == "53.8%"
    assert T.num(-6.97435, 3) == "-6.974"
    assert T.count(4141.0) == "4,141"


def _board():
    rows = []
    for m, q, b, acc in (("gbm", -6.9, 0.2497, 0.538), ("climatology", -6.6, 0.2499, 0.518),
                         ("static_gnn", -7.0, 0.2505, 0.504)):
        for h in (1, 5):
            rows.append({"model": m, "h": h, "vol_qlike": q + h, "vol_rmse_vol": 0.01 * h, "vol_mz_r2": 0.1,
                         "brier": b, "logloss": 0.69, "auc": 0.5, "accuracy": acc, "balanced_accuracy": 0.5,
                         "up_calls": 0.6, "up_precision": 0.52, "down_precision": np.nan if m == "climatology" else 0.5})
    return pd.DataFrame(rows)


def test_boards_follow_the_ladder_and_mark_the_best():
    f, bold = T.volatility_board(_board())
    assert f["Model"].tolist() == ["climatology", "gbm", "static_gnn"]
    assert (2, "QLIKE (1-day)") in bold and (0, "QLIKE (1-day)") not in bold  # lowest QLIKE is best
    assert {(i, "RMSE (1-day)") for i in range(3)} <= bold  # ties are all marked
    f, bold = T.direction_board(_board(), 1)
    assert (1, "Brier score") in bold and (1, "Accuracy") in bold
    assert f.loc[0, "Down calls correct"] == "–"
    assert not any(c == "Share called up" for _, c in bold)  # a share, not a score


def test_significance_table_puts_horizons_side_by_side():
    sig = pd.DataFrame([{"model": m, "vs": "climatology", "loss": "qlike", "h": h, "mean_diff": -0.3 * h,
                         "dm_stat": -5.0, "dm_p": 1e-5, "wilcoxon_p": 0.04, "perm_p": 0.2, "folds_better": "7/8"}
                        for m in ("gbm", "linear") for h in (1, 5)])
    t = T.significance_table(sig, "qlike")
    assert t["Model"].tolist() == ["linear", "gbm"]
    assert t.loc[0, "Mean diff. (5-day)"] == "-1.500" and t.loc[0, "p DM (1-day)"] == "<0.001*" and t.loc[0, "p permutation (5-day)"] == "0.200"


def test_markdown_and_files(tmp_path):
    f, bold = T.volatility_board(_board())
    t = T.Table("T01-demo", "Demo", f, "A note.", ("x/y.csv",), bold=bold)
    md = T.to_markdown(t)
    assert md.splitlines()[0].startswith("| Model | QLIKE (1-day)") and "**-6.000**" in md
    written = T.write_pack([t], tmp_path)
    assert {p.name for p in written} == {"T01-demo.csv", "T01-demo.md", "index.md", "tables.docx"}
    assert pd.read_csv(tmp_path / "T01-demo.csv", dtype=str).equals(f.astype(str))
    from docx import Document

    doc = Document(tmp_path / "tables.docx")
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 4
    assert doc.tables[0].cell(3, 1).text == "-6.000"


@pytest.mark.skipif(not (REPORTS_DIR / "phase4-aimdg" / "scoreboard.csv").exists(), reason="no reports")
def test_pack_is_faithful_to_the_reports():
    pack = {t.key: t for t in T.build(REPORTS_DIR)}
    assert len(pack) == 27 and all(len(t.frame) and t.sources for t in pack.values())
    raw = pd.read_csv(REPORTS_DIR / "phase4-aimdg" / "scoreboard.csv").set_index(["model", "h"])
    t = pack["T01-volatility-scoreboard"].frame.set_index("Model")
    assert t.index.tolist() == list(T.LADDER)
    for m in T.LADDER:
        assert t.loc[m, "QLIKE (5-day)"] == f"{raw.loc[(m, 5), 'vol_qlike']:.3f}"
    sig = pd.read_csv(REPORTS_DIR / "phase4-aimdg-vs-static_gnn" / "significance.csv")
    r = sig[(sig.model == "aimdg") & (sig.loss == "qlike") & (sig.h == 1)].iloc[0]
    assert pack["T07-aimdg-vs-static-gnn"].frame.loc[0, "p (DM)"] == T.pval(r.dm_p)
